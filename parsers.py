import io
import os
import json
import tempfile
import pandas as pd
import pypdf
import docx2txt
import openpyxl

def extract_pdf(file_bytes: bytes) -> dict:
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    num_pages = len(reader.pages)
    sections = []
    full_text_parts = []
    
    for i, page in enumerate(reader.pages):
        page_num = i + 1
        page_text = page.extract_text() or ""
        page_text_clean = page_text.strip()
        if page_text_clean:
            sections.append({
                "source_ref": f"Page {page_num}",
                "content": page_text_clean,
                "page": page_num
            })
            full_text_parts.append(f"--- [Page {page_num}] ---\n{page_text_clean}")
            
    full_text = "\n\n".join(full_text_parts)
    return {
        "text": full_text,
        "sections": sections,
        "metadata": {
            "type": "PDF Document",
            "page_count": num_pages,
            "char_count": len(full_text)
        }
    }

def extract_docx(file_bytes: bytes) -> dict:
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(file_bytes)
        tmp_path = tmp.name
        
    try:
        text = docx2txt.process(tmp_path)
    finally:
        if os.path.exists(tmp_path):
            try:
                os.remove(tmp_path)
            except Exception:
                pass
                
    text = (text or "").strip()
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    sections = []
    
    for i, p in enumerate(paragraphs):
        sections.append({
            "source_ref": f"Section {i + 1}",
            "content": p
        })
        
    return {
        "text": text,
        "sections": sections,
        "metadata": {
            "type": "Word Document (.docx)",
            "paragraph_count": len(paragraphs),
            "char_count": len(text)
        }
    }

def extract_excel(file_bytes: bytes, filename: str) -> dict:
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheet_names = wb.sheetnames
    sections = []
    full_text_parts = []
    total_rows = 0

    for sheet_name in sheet_names:
        sheet = wb[sheet_name]
        data = []
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None and str(cell).strip() != "" for cell in row):
                data.append([str(c) if c is not None else "" for c in row])

        if data:
            total_rows += len(data)
            # Format as readable markdown table or lines
            header = data[0]
            rows = data[1:]
            sheet_text_lines = [f"### Sheet: {sheet_name}"]
            sheet_text_lines.append("Columns: " + " | ".join(header))
            sheet_text_lines.append("-" * 40)
            for r in rows[:200]:  # Cap individual sheet view rows to keep reasonable
                sheet_text_lines.append(" | ".join(r))
            
            sheet_content = "\n".join(sheet_text_lines)
            sections.append({
                "source_ref": f"Sheet: {sheet_name}",
                "content": sheet_content,
                "sheet": sheet_name,
                "row_count": len(data)
            })
            full_text_parts.append(sheet_content)

    full_text = "\n\n".join(full_text_parts)
    is_macro = filename.lower().endswith(".xlsm")
    doc_type = "Excel Macro-Enabled Workbook (.xlsm)" if is_macro else "Excel Workbook (.xlsx)"
    
    return {
        "text": full_text,
        "sections": sections,
        "metadata": {
            "type": doc_type,
            "sheets": sheet_names,
            "sheet_count": len(sheet_names),
            "total_rows": total_rows,
            "char_count": len(full_text)
        }
    }

def extract_csv(file_bytes: bytes) -> dict:
    for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            content = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        content = file_bytes.decode("utf-8", errors="ignore")

    df = pd.read_csv(io.StringIO(content))
    summary_lines = [
        f"Columns ({len(df.columns)}): {', '.join([str(c) for c in df.columns])}",
        f"Total Rows: {len(df)}",
        "\nData Preview (first 100 rows):"
    ]
    summary_lines.append(df.head(100).to_string(index=False))
    text = "\n".join(summary_lines)

    return {
        "text": text,
        "sections": [{"source_ref": "CSV Table", "content": text}],
        "metadata": {
            "type": "CSV Spreadsheet",
            "rows": len(df),
            "columns": len(df.columns),
            "char_count": len(text)
        }
    }

def extract_plain_text(file_bytes: bytes, filename: str) -> dict:
    for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
        try:
            text = file_bytes.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = file_bytes.decode("utf-8", errors="ignore")

    ext = filename.split(".")[-1].lower() if "." in filename else "txt"
    type_map = {
        "md": "Markdown Document",
        "json": "JSON Data",
        "html": "HTML Document",
        "htm": "HTML Document",
        "py": "Python Script",
        "js": "JavaScript File",
        "txt": "Plain Text File"
    }

    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    sections = [{"source_ref": f"Section {i+1}", "content": p} for i, p in enumerate(paragraphs)]
    if not sections:
        sections = [{"source_ref": "Full Document", "content": text}]

    return {
        "text": text,
        "sections": sections,
        "metadata": {
            "type": type_map.get(ext, f"{ext.upper()} File"),
            "char_count": len(text),
            "paragraph_count": len(paragraphs)
        }
    }

def parse_document(file_bytes: bytes, filename: str) -> dict:
    """Dispatches to the appropriate parser based on file extension."""
    ext = filename.split(".")[-1].lower() if "." in filename else ""
    
    if ext == "pdf":
        return extract_pdf(file_bytes)
    elif ext in ["docx", "doc"]:
        return extract_docx(file_bytes)
    elif ext in ["xlsx", "xlsm", "xls"]:
        return extract_excel(file_bytes, filename)
    elif ext == "csv":
        return extract_csv(file_bytes)
    else:
        return extract_plain_text(file_bytes, filename)
